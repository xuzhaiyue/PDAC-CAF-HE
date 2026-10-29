"""Supplementary tables v2: combine the TSVs written by the prepare script into one xlsx (README sheet + one sheet per table), then read back with openpyxl and check every cell.
scRNA environment (xlsxwriter to write, openpyxl to read); nice 19.
"""
from __future__ import annotations

import json
import math
import os
import sys
from pathlib import Path

import openpyxl
import pandas as pd

PROJ = Path(__file__).resolve().parents[1]
OUT = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else PROJ / "results/tables/2026-10-28_run02"
XLSX = OUT / "Supplementary_Tables.xlsx"
CONVENTIONS = [
    "Static transcription of frozen analysis outputs; no analysis was refitted or retested for these tables.",
    "Blank cells mean not recorded or not applicable, never zero. 'not_recorded' treatment is not confirmed untreated.",
    "S0 = TCGA bulk-supervised model; S1 = pan-cancer spatial model (156 non-pancreatic HEST sections); "
    "S2 = PDAC spatial fine-tuning, leave one patient out.",
    "Local ρ = Spearman correlation between predicted and measured programme-minus-reference index over CAF-rich regions of one section.",
    "Scale: public GEO Visium µm were originally derived assuming a 55-µm spot; Space Ranger spot_diameter_fullres spans 65 µm. "
    "S6 um_per_px is corrected (legacy × 65/55); nominal 112/128-µm GEO fields were ~132/151 µm. ρ, regions and masks do not depend on this.",
    "The independent 10x Visium HD section comes from the same tissue source as HEST TENX179 and must not be counted as a second independent sample.",
    "Amendments A10–A16 were added after A9 failed its section-level criterion; each was frozen in config/fig1.yaml before it was run.",
]


def main():
    if os.getpriority(os.PRIO_PROCESS, 0) < 19:
        os.nice(19 - os.getpriority(os.PRIO_PROCESS, 0))
    assert not XLSX.exists(), "do not overwrite existing results"
    index = pd.read_csv(OUT / "README_index.tsv", sep="\t")
    sheets = json.loads((OUT / "sheets.json").read_text()); assert sheets == index.sheet.tolist()
    tables = {s: pd.read_csv(OUT / f"{s}.tsv", sep="\t", keep_default_na=True) for s in sheets}

    with pd.ExcelWriter(XLSX, engine="xlsxwriter") as xw:
        wb = xw.book
        bold = wb.add_format({"bold": True, "font_name": "Arial", "font_size": 10})
        head = wb.add_format({"bold": True, "font_name": "Arial", "font_size": 9, "bg_color": "#E8EDF1", "bottom": 1, "text_wrap": True, "valign": "top"})
        body = wb.add_format({"font_name": "Arial", "font_size": 9})
        wrap = wb.add_format({"font_name": "Arial", "font_size": 9, "text_wrap": True, "valign": "top"})
        ws = wb.add_worksheet("README")
        ws.write(0, 0, "Supplementary Tables — CTHRC1+ matrix CAF programme read from H&E (v2, 2026-10-28)", bold)
        r = 2
        for line in CONVENTIONS:
            ws.write(r, 0, "•", body); ws.merge_range(r, 1, r, 6, line, wrap); ws.set_row(r, 26); r += 1
        r += 1
        for j, c in enumerate(["Sheet", "Title", "Figure", "Rows", "Source", "Note"]):
            ws.write(r, j + 1, c, head)
        for k, row in index.iterrows():
            r += 1
            for j, c in enumerate(["sheet", "title", "figure", "n_rows", "source", "note"]):
                v = row[c]; ws.write(r, j + 1, "" if pd.isna(v) else v, wrap)
            if k < len(index) and isinstance(row.sheet, str):
                ws.write_url(r, 1, f"internal:'{row.sheet}'!A1", string=row.sheet)
        for j, w in enumerate([3, 20, 44, 22, 6, 42, 70]):
            ws.set_column(j, j, w)
        for s in sheets:
            df = tables[s]; df.to_excel(xw, sheet_name=s, index=False)
            w = xw.sheets[s]
            for j, c in enumerate(df.columns):
                w.write(0, j, c, head)
                width = min(48, max(9, len(str(c)) * .9, df[c].astype(str).str.len().quantile(.9) * .9 if len(df) else 9))
                w.set_column(j, j, width, body)
            w.set_row(0, 30); w.freeze_panes(1, 1); w.autofilter(0, 0, max(len(df), 1), len(df.columns) - 1)

    # ---- Read back and check every cell
    book = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    assert book.sheetnames == ["README"] + sheets
    n_cells, errors = 0, []
    for s in sheets:
        df = tables[s]; rows = list(book[s].iter_rows(values_only=True))
        assert list(rows[0][:len(df.columns)]) == list(df.columns), s
        assert len(rows) - 1 == len(df), (s, len(rows) - 1, len(df))
        for i, vals in enumerate(rows[1:]):
            for j, c in enumerate(df.columns):
                a, b = df.iat[i, j], vals[j]; n_cells += 1
                if pd.isna(a):
                    ok = b is None or b == ""
                elif isinstance(a, (bool,)) or str(a) in ("True", "False"):
                    ok = str(a) == str(b) or bool(a) == b
                elif isinstance(a, (int, float)) and not isinstance(a, bool):
                    ok = isinstance(b, (int, float)) and math.isclose(float(a), float(b), rel_tol=1e-12, abs_tol=1e-300)
                else:
                    ok = str(a) == str(b)
                if not ok:
                    errors.append((s, i, c, repr(a), repr(b)))
    book.close()
    res = {"xlsx": str(XLSX.relative_to(PROJ)), "sheets": len(sheets), "data_rows": int(index.n_rows.sum()), "cells_checked": n_cells,
           "mismatches": len(errors), "first_mismatches": errors[:10]}
    (OUT / "verification.json").write_text(json.dumps(res, indent=2, default=str)); print(json.dumps(res, indent=1, default=str))
    assert not errors


if __name__ == "__main__":
    main()
